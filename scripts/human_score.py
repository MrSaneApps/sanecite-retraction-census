#!/usr/bin/env python3
"""Score the human labelling round (PROTOCOL.md, Validation; rules fixed in amendment 4 before any answers exist).

Each labeller's two answers per passage become a label through classify.label(), the same rule the models' answers
go through, with the same citation-group count from citegroup.py. A citing paper's label is the strongest across its
passages. Reports: completeness; agreement and Cohen's kappa between the two people; the panel's agreement with each
person and with the papers both people label the same; the same on the headline question (substantive or not); and
whether the 85% rule is triggered (panel agreement below 85% with EITHER labeller).

usage: SANECITE_DATA=<data dir> human_score.py <packet_dir> <labeller A .xlsx|.csv> <labeller B .xlsx|.csv> [out.json]"""
import collections, csv, json, os, sys
sys.path.insert(0, os.path.dirname(__file__)); import classify as C

DATA, PACKET, FA, FB = os.environ["SANECITE_DATA"], sys.argv[1], sys.argv[2], sys.argv[3]
OUT = sys.argv[4] if len(sys.argv) > 4 else None
RANK = {"aware": 3, "substantive": 2, "perfunctory": 1, "unclear": 0}
Q2 = {"specific", "counted", "background", "described", "unclear"}
RULE = 0.85

draw = json.load(open(os.path.join(PACKET, "draw.json")))
item_pair = {it["item"]: it["pair_id"] for it in draw["items"]}
want = set(item_pair.values())
recs = {}
for l in open(os.path.join(DATA, "stage2_v31.jsonl")):
    try: r = json.loads(l)
    except Exception: continue
    if r.get("pair_id") in want: recs[r["pair_id"]] = r
panel = {}
for l in open(os.path.join(DATA, "census_rows.jsonl")):
    r = json.loads(l)
    if r["pair_id"] in want: panel[r["pair_id"]] = r["label"]


def read(path):
    """Rows of (paper, passage, sentence, q1, q2) from a returned workbook or csv."""
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)["Passages"]
        it = ws.iter_rows(min_row=2, values_only=True)
        return [(r[0], r[1], r[5], r[7], r[8]) for r in it if r[0] is not None]
    rows = list(csv.reader(open(path, encoding="utf-8")))[1:]
    return [(r[0], r[1], r[6], r[8], r[9]) for r in rows if r and r[0]]


def labels(path):
    """Pair labels for one labeller, plus problems found (missing, invalid, or rows that do not match the packet)."""
    by_pair, problems = collections.defaultdict(dict), []
    for paper, passage, sentence, q1, q2 in read(path):
        k, i = int(paper), int(passage) - 1
        pid = item_pair.get(k)
        if pid is None or i >= len(recs[pid]["contexts"]):
            problems.append(f"paper {paper} passage {passage}: not in the packet"); continue
        if (sentence or "").strip() != recs[pid]["contexts"][i]["sentence"].strip():
            problems.append(f"paper {paper} passage {passage}: sentence does not match the packet (rows edited or moved?)"); continue
        q1, q2 = str(q1 or "").strip().lower(), str(q2 or "").strip().lower()
        if q1 not in ("yes", "no") or q2 not in Q2:
            problems.append(f"paper {paper} passage {passage}: missing or invalid answer ({q1!r}, {q2!r})"); continue
        rec = recs[pid]
        text = " ".join(C.passage(rec, i)).strip()
        by_pair[pid][i] = C.label("disputed" if q1 == "yes" else "neutral", q2, C.group_n(rec, i), text)
    out = {}
    for pid, ctx in by_pair.items():
        if len(ctx) == len(recs[pid]["contexts"]):  # a paper counts only when every passage is answered
            out[pid] = max(ctx.values(), key=lambda x: RANK[x])
    return out, problems


def kappa(a, b, keys):
    n = len(keys)
    if not n: return None
    po = sum(a[k] == b[k] for k in keys) / n
    ca, cb = collections.Counter(a[k] for k in keys), collections.Counter(b[k] for k in keys)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n / n
    return round((po - pe) / (1 - pe), 3) if pe < 1 else None


def agree(a, b, keys, binary=False):
    f = (lambda x: x == "substantive") if binary else (lambda x: x)
    n = len(keys); ok = sum(f(a[k]) == f(b[k]) for k in keys)
    return {"agree": ok, "n": n, "rate": round(ok / n, 4) if n else None}


def confusion(a, b, keys):
    m = collections.Counter((a[k], b[k]) for k in keys)
    return {f"{x}->{y}": c for (x, y), c in sorted(m.items())}


A, pa = labels(FA)
B, pb = labels(FB)
both = sorted(set(A) & set(B))
agreed = [k for k in both if A[k] == B[k]]
consensus = {k: A[k] for k in agreed}
res = {
    "packet": {"papers": len(item_pair), "passages": sum(len(recs[p]["contexts"]) for p in want)},
    "complete_papers": {"A": len(A), "B": len(B), "both": len(both)},
    "problems": {"A": pa, "B": pb},
    "human_vs_human": {**agree(A, B, both), "kappa": kappa(A, B, both), "substantive_or_not": agree(A, B, both, True),
                       "confusion_A_to_B": confusion(A, B, both)},
    "panel_vs_A": {**agree(panel, A, sorted(A)), "kappa": kappa(panel, A, sorted(A)),
                   "substantive_or_not": agree(panel, A, sorted(A), True), "confusion_panel_to_A": confusion(panel, A, sorted(A))},
    "panel_vs_B": {**agree(panel, B, sorted(B)), "kappa": kappa(panel, B, sorted(B)),
                   "substantive_or_not": agree(panel, B, sorted(B), True), "confusion_panel_to_B": confusion(panel, B, sorted(B))},
    "panel_vs_consensus": {**agree(panel, consensus, agreed), "kappa": kappa(panel, consensus, agreed),
                           "substantive_or_not": agree(panel, consensus, agreed, True),
                           "confusion_panel_to_consensus": confusion(panel, consensus, agreed)},
}
low = [w for w in ("A", "B") if (res[f"panel_vs_{w}"]["rate"] or 0) < RULE]
res["rule_85"] = {"triggered": bool(low), "below_85_with": low,
                  "meaning": "headline must be reported with this agreement and corrected with the human labels" if low
                  else "panel agreement is at least 85% with both labellers"}
if pa or pb or len(both) < len(item_pair):
    res["warning"] = "incomplete or problem rows: fix them with the labeller before treating these numbers as final"
print(json.dumps(res, indent=1, ensure_ascii=False))
if OUT: json.dump(res, open(OUT, "w"), indent=1, ensure_ascii=False)
