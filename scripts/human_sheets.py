#!/usr/bin/env python3
"""Turn the human labelling packet (passages.csv from human_packet.py) into one Excel workbook per labeller:
a read-me sheet with the instructions and a passages sheet with drop-down answers. Same passages, same order, for
both labellers; nothing in the workbook shows a model label.
usage: human_sheets.py <packet_dir> <out_dir> [labeller ids, default A B]"""
import csv, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

PACKET, OUT = sys.argv[1], sys.argv[2]
WHO = sys.argv[3:] or ["A", "B"]
rows = list(csv.reader(open(os.path.join(PACKET, "passages.csv"), encoding="utf-8")))[1:]
instructions = open(os.path.join(PACKET, "INSTRUCTIONS.md"), encoding="utf-8").read()

COLS = [("Paper", 7), ("Passage", 8), ("Cited work (the target)", 30), ("Section", 16), ("Text before", 42),
        ("Citing sentence (→ marks the target)", 55), ("Text after", 42),
        ("Q1: says anything against the target?", 14), ("Q2: how is the target used?", 15), ("Notes (optional)", 28)]
Q1_COL, Q2_COL = "H", "I"
WRAP = Alignment(wrap_text=True, vertical="top")
HEAD = PatternFill("solid", fgColor="1F3A5F")
SHADE = PatternFill("solid", fgColor="EEF3F8")
ANSWER = PatternFill("solid", fgColor="FFF8DC")
THIN = Border(bottom=Side(style="thin", color="C8D2DC"))


def plain(md):
    """Markdown to readable sheet lines: drop # ** and backticks, join hard-wrapped lines into one paragraph,
    and point at the Passages tab instead of the csv."""
    out = []
    for line in md.replace("in `passages.csv`", "in the Passages tab").splitlines():
        line = re.sub(r"^#+\s*", "", line).replace("**", "").replace("`", "").rstrip()
        starts_new = not line or line.startswith(("- ", "Q1_", "Q2_", "Examples", "Labelling"))
        if out and out[-1] and not starts_new and not out[-1].endswith(":"):
            out[-1] += " " + line
        else:
            out.append(line)
    return out


def build(who):
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me first"
    rd.column_dimensions["A"].width = 120
    intro = [f"SaneCite labelling sheet, labeller {who}", "",
             "How to fill this in:",
             "1. Read the instructions below once, including the examples.",
             "2. Go to the 'Passages' tab. For every row, pick an answer from the drop-down in the two yellow columns (Q1 and Q2).",
             "3. Rows with the same Paper number come from the same citing paper. Answer each row on its own.",
             "4. Save the file and send it back as it is. Do not sort, delete or reorder rows.",
             f"Rows answered so far (of {len(rows)}):", ""]
    for i, line in enumerate(intro + plain(instructions), 1):
        c = rd.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1: c.font = Font(bold=True, size=14)
        elif line.endswith(":") or line.startswith(("Q1_", "Q2_", "Examples", "Labelling")): c.font = Font(bold=True)
    n = len(rows)
    rd["B8"] = f'=SUMPRODUCT((Passages!{Q1_COL}2:{Q1_COL}{n + 1}<>"")*(Passages!{Q2_COL}2:{Q2_COL}{n + 1}<>""))'
    rd["B8"].font = Font(bold=True)
    rd.column_dimensions["B"].width = 10

    ws = wb.create_sheet("Passages")
    for j, (name, width) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 34
    last_item, shade = None, False
    for i, r in enumerate(rows, 2):
        item, passage, title, section, where, before, sentence, after = r[:8]
        if item != last_item: shade, last_item = not shade, item
        sec = (section or "") + (" (table row)" if where == "table" else "")
        vals = [int(item), int(passage), title, sec, before, sentence, after, None, None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = WRAP
            c.border = THIN
            if j in (8, 9): c.fill = ANSWER
            elif shade: c.fill = SHADE
        longest = max(len(before) / 42, len(sentence) / 55, len(after) / 42, len(title) / 30, 1)
        ws.row_dimensions[i].height = min(409, 15 * (int(longest) + 2))
    last = len(rows) + 1
    q1 = DataValidation(type="list", formula1='"yes,no"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Q1", error="Choose yes or no.")
    q2 = DataValidation(type="list", formula1='"specific,counted,background,described,unclear"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Q2", error="Choose one of the five answers.")
    ws.add_data_validation(q1); ws.add_data_validation(q2)
    q1.add(f"{Q1_COL}2:{Q1_COL}{last}"); q2.add(f"{Q2_COL}2:{Q2_COL}{last}")
    # answered rows turn green, so a skipped row stands out
    green = PatternFill("solid", fgColor="DDF2DD")
    ws.conditional_formatting.add(f"{Q1_COL}2:{Q2_COL}{last}", FormulaRule(formula=[f'AND(${Q1_COL}2<>"",${Q2_COL}2<>"")'], fill=green))
    ws.freeze_panes = "C2"
    wb.active = 0
    path = os.path.join(OUT, f"SaneCite_labelling_{who}.xlsx")
    wb.save(path)
    return path


os.makedirs(OUT, exist_ok=True)
for w in WHO: print(build(w))
